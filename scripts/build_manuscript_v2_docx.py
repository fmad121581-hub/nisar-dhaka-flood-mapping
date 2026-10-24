"""
Build the v2 manuscript (.docx) from manuscript_v2/manuscript_v2.md, using the earlier NHESS
submission as the Word template (its styles, page set-up and line numbering are kept; its
body is emptied). References come from manuscript_v2/05_references.md.

Markup understood in manuscript_v2.md:
  TITLE:/AUTHORS:/AFFILIATION:/CORRESPONDENCE:/KEYWORDS:/ABSTRACT:  front matter
  # / ##                         Heading 1 / Heading 2
  TABLE: caption  + "| a | b |" rows (first row = header)
  BACKMATTER: heading  + one paragraph
  FIGURE: path | caption         placed after the section whose number is in FIG_AFTER
Output: manuscript_v2/Ahmed_NISAR_Bangladesh_v2_NHESS.docx
"""
from pathlib import Path
import copy
import re
import docx
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(r"D:\NISAR_Dhaka_Flood")
TEMPLATE = Path(r"C:\Users\user\Downloads\Ahmed_NISAR_Dhaka_Manuscript_NHESS_CopernicusTemplate.docx")
SRC = ROOT / "manuscript_v2" / "manuscript_v2.md"
REFS = ROOT / "manuscript_v2" / "05_references.md"
OUT = ROOT / "manuscript_v2" / "Ahmed_NISAR_Bangladesh_v2_NHESS.docx"
# figure number -> heading after whose section it is inserted
FIG_AFTER = {1: "2.2 SAR data", 2: "3.1 Granule screening",
             3: "3.2 Hakaluki Haor through the recession", 4: "3.3 Kaptai Lake during the reservoir rise",
             5: "3.4 Dhaka during the July pluvial event"}

doc = docx.Document(TEMPLATE)
body = doc.element.body
for el in list(body):
    if el.tag != qn("w:sectPr"):
        body.remove(el)


def para(text, style="Normal", bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        text = text[len(bold_prefix):]
    p.add_run(text)
    return p


def rule(cell, edge):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    b = OxmlElement(f"w:{edge}")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "0"), ("w:color", "000000")):
        b.set(qn(k), v)
    borders.append(b)


def table(caption, rows):
    para(caption, "Caption")
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = val
            for r in c.paragraphs[0].runs:
                r.font.size = Pt(8)
                r.bold = i == 0
            if i == 0:
                rule(c, "top")
                rule(c, "bottom")
            if i == len(rows) - 1:
                rule(c, "bottom")
    doc.add_paragraph()


def figure(path, caption):
    p = doc.add_paragraph()
    p.add_run().add_picture(str(ROOT / path), width=Cm(15.5))
    para(caption, "Caption")


lines = SRC.read_text(encoding="utf-8").splitlines()
figs = {}
for ln in lines:
    if ln.startswith("FIGURE:"):
        path, cap = [s.strip() for s in ln[7:].split("|", 1)]
        figs[int(re.search(r"Figure (\d+)", cap).group(1))] = (path, cap)
pending = {v: k for k, v in FIG_AFTER.items()}
current = None


def flush_fig():
    if current in pending:
        figure(*figs[pending.pop(current)])


i = 0
backmatter_started = False
while i < len(lines):
    ln = lines[i].rstrip()
    if not ln or ln.startswith("FIGURE:"):
        i += 1
        continue
    if ln.startswith("TITLE:"):
        para(ln[6:].strip(), "MS title")
    elif ln.startswith("AUTHORS:"):
        para(ln[8:].strip(), "Authors")
    elif ln.startswith("AFFILIATION:"):
        para(ln[12:].strip(), "Affiliation")
    elif ln.startswith("CORRESPONDENCE:"):
        para(ln[15:].strip(), "Correspondence")
    elif ln.startswith("KEYWORDS:"):
        para(ln[9:].strip(), bold_prefix="Keywords:")
    elif ln.startswith("ABSTRACT:"):
        para("Abstract. " + ln[9:].strip(), bold_prefix="Abstract.")
    elif ln.startswith("## "):
        flush_fig()
        current = ln[3:].strip()
        para(current, "Heading 2")
    elif ln.startswith("# "):
        flush_fig()
        current = ln[2:].strip()
        para(current, "Heading 1")
    elif ln.startswith("TABLE:"):
        cap = ln[6:].strip()
        rows = []
        i += 1
        while i < len(lines) and lines[i].startswith("|"):
            rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        table(cap, rows)
        continue
    elif ln.startswith("BACKMATTER:"):
        if not backmatter_started:
            flush_fig()
            current = None
            backmatter_started = True
        para(ln[11:].strip(), "Heading 1")
        para(lines[i + 1].strip())
        i += 2
        continue
    else:
        para(ln)
    i += 1
for n in sorted(pending.values()):  # anything not yet placed
    figure(*figs[n])

para("References", "Heading 1")
for ln in REFS.read_text(encoding="utf-8").split("## Not cited")[0].splitlines():
    if ln and not ln.startswith("#"):
        para(re.sub(r"\s*\[[^\]]*\]\s*$", "", ln))
doc.save(OUT)
print(f"Wrote {OUT}")

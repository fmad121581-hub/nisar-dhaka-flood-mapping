"""
Build the separate supplement file (.docx) for manuscript v2 from the CSV tables written by
make_supplement_tables.py, using the earlier submission as the Word template.
Output: manuscript_v2/Ahmed_NISAR_Bangladesh_v2_supplement.docx
"""
from pathlib import Path
import csv
import docx
from docx.shared import Pt, Cm
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

ROOT = Path(r"D:\NISAR_Dhaka_Flood")
SUP = ROOT / "manuscript_v2" / "supplement"
TEMPLATE = Path(r"C:\Users\user\Downloads\Ahmed_NISAR_Dhaka_Manuscript_NHESS_CopernicusTemplate.docx")
OUT = ROOT / "manuscript_v2" / "Ahmed_NISAR_Bangladesh_v2_supplement.docx"
TABLES = [
    ("tableS1_granules", "Table S1. NISAR GCOV granules, OPERA RTC-S1 bursts and HyP3 RTC scenes used, with the screening outcome. HHHH_mean_value is the layer attribute stored in each GCOV granule."),
    ("tableS2_stage_and_levels", "Table S2. River-stage (Fenchuganj, Kushiyara; danger level 9.45 m MSL in 2026 bulletins) and Kaptai Lake level records, with the date of the reading, its quality and its source. Several Fenchuganj values come from news reports of FFWC bulletins or from a secondary compilation and lie 1–7 days from the NISAR date."),
    ("tableS3_nesz_dhaka", "Table S3. Noise-equivalent backscatter (median over each ward and a 2 km margin) from the GCOV calibration metadata of the four Dhaka granules."),
    ("tableS4_hakaluki_basin_vs_hill", "Table S4. Hakaluki tree cover in the basin (Copernicus surface model below 12 m), at 12–25 m and on the hills (above 40 m): class-mean HH and HV, HH−HV and the share of pixels above −5.9 dB HH, per NISAR date."),
    ("tableS5_class_means_ci", "Table S5. Class means of co-polarized backscatter with 95 % spatial block-bootstrap intervals, co- to cross-polarized ratio and water-like share on each sensor's pooled Otsu threshold, Hakaluki and Kaptai."),
]

doc = docx.Document(TEMPLATE)
body = doc.element.body
for el in list(body):
    if el.tag != qn("w:sectPr"):
        body.remove(el)
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width

doc.add_paragraph("Supplement to: Early NISAR L-band observations of monsoon flooding in Bangladesh, "
                  "compared with Sentinel-1 against independent water-level records", style="MS title")
doc.add_paragraph("Fahim Ahmed", style="Authors")
for name, caption in TABLES:
    rows = list(csv.reader(open(SUP / f"{name}.csv", encoding="utf-8")))
    doc.add_paragraph(caption, style="Caption")
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = doc.styles["Normal Table"]
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = val
            for r in c.paragraphs[0].runs:
                r.font.size = Pt(6.5)
                r.bold = i == 0
    doc.add_paragraph()
doc.add_paragraph("Machine-readable copies of Tables S1–S5 (CSV) are archived with the code release.")
doc.save(OUT)
print(f"Wrote {OUT}")
